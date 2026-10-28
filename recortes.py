import os
import glob
import re
import socket
import datetime
import openpyxl

IP_ZEBRA = "192.168.1.95"
PUERTO = 9100

# Diccionario maestro de traducción y contracciones de acabados
CATALOGO_FACES = {
    # Variantes especiales contraídas
    "DOLPHIN_GREY_VINYL2D": "D.G._VINYL2D",
    "DG_VINYL2D": "D.G._VINYL2D",
    "DG-V2D": "D.G._VINYL2D",
    "D.G._VINYL2D": "D.G._VINYL2D",
    "PERFECT_WHITE_VINYL2D": "PW-V2D",
    "PW_VINYL2D": "PW-V2D",
    "PW-V2D": "PW-V2D",
    "CHARLESTON_GREY_VINYL2D": "CG-V2D",
    "CG_VINYL2D": "CG-V2D",
    "CG-V2D": "CG-V2D",

    # Acabados estándar
    # Acabados estándar
    "W": "PERFECT_WHITE",
    "PW": "PERFECT_WHITE",
    "CH.G.": "CHARLESTON_GREY",
    "DG": "DOLPHIN_GREY",
    "FG": "FOSSIL_GREY",
    "GBR": "GIBRALTAR",
    "GB": "GIBRALTAR",
    "HM": "HONEY_MAPLE",
    "HMS": "HONEY_MAPLE",
    "SOM": "SOMERSET",
    "NOK": "NATURAL_OAK",
    "JV": "JAVA",
    "CAR": "CARMEL",
    "CNN": "CINNAMON",
    "VB": "VELVET_BLUE",
    "VTN": "VENETIAN",
    "VENETIAN": "VENETIAN",
    "W_61": "WHITE_61",
    "WHITE_61": "WHITE_61",
    "M": "MAPLE",
    "MP": "MAPLE",
    "BK": "BLACK",
    "SLT": "SLATE",
    "RAW": "RAW_MDF",
    "PB": "PARTICLE_BOARD"
    
}

def obtener_archivo_reciente():
    """Localiza el archivo Excel más reciente en record/ o en la carpeta raíz."""
    rutas = glob.glob(os.path.join("record", "*.xlsx"))
    if not rutas:
        rutas = [f for f in glob.glob("*.xlsx") if not f.startswith("~$")]
    if not rutas:
        return None
    rutas.sort(key=os.path.getmtime, reverse=True)
    return rutas[0]

def formatear_fecha_mes_dia(texto_lote):
    """Extrae únicamente Mes y Día (MM  -  DD) ignorando prefijos de corrida."""
    texto = str(texto_lote).strip()

    match_slash = re.search(r'(?<!\d)(\d{1,2})/(\d{1,2})', texto)
    if match_slash:
        m, d = map(int, match_slash.groups())
        if 1 <= m <= 12 and 1 <= d <= 31:
            return f"{m:02d}  -  {d:02d}"

    texto_limpio = re.sub(r'\b[A-Za-z]+\d*[-_]', '', texto)

    match_dash = re.search(r'(?<![A-Za-z0-9])(\d{1,2})-(\d{1,2})', texto_limpio)
    if match_dash:
        m, d = map(int, match_dash.groups())
        if 1 <= m <= 12 and 1 <= d <= 31:
            return f"{m:02d}  -  {d:02d}"

    hoy = datetime.datetime.now()
    return f"{hoy.month:02d}  -  {hoy.day:02d}"

def extraer_color_lote(texto_meta):
    """Detecta el color de la orden."""
    texto = str(texto_meta).upper()
    colores = [
        "PINK", "ROSA", "BLUE", "AZUL", "GREEN", "VERDE", 
        "ORANGE", "NARANJA", "PURPLE", "MORADO", "LAVANDER", "LAVANDA", "RED", "ROJO"
    ]
    for c in colores:
        if c in texto:
            return c
    palabras = texto.split()
    if len(palabras) > 1 and not any(ch.isdigit() for ch in palabras[-1]):
        return palabras[-1]
    return "RECUT"

def generar_cadena_marca_agua(color):
    """Control de repetición discreto según longitud del color."""
    c = str(color).strip().upper()
    if c in ["PINK", "ROSA", "BLUE", "AZUL"]:
        return f"{c}   {c}   {c}   {c}"
    elif c in ["GREEN", "VERDE"]:
        return f"{c}   {c}   {c}"
    elif c in ["LAVANDER", "LAVANDA", "ORANGE", "NARANJA", "PURPLE", "MORADO"]:
        return f"{c}   {c}"
    else:
        return f"{c}   {c}"

def traducir_acabado(sigla):
    s = sigla.strip().upper()
    return CATALOGO_FACES.get(s, s)

def desglosar_faces_completo(faces_val):
    val = str(faces_val or "").strip()
    if not val or val.upper() in ["NONE", "N/A", "NAN"]:
        return "N/A", "N/A"
    if "/" in val:
        partes = val.split("/", 1)
        return traducir_acabado(partes[0]), traducir_acabado(partes[1])
    res = traducir_acabado(val)
    return res, res

def generar_zpl_maestro(fecha, part_name, l, w, t, qty, routing, face1, face2, es_amarilla, color_watermark=""):
    """
    Generador ZPL maestro unificado:
    - Estandarización de Face 1, Face 2 y cantos a 30, 34.
    - Cápsula Length + GRAIN.
    - Pie de página condicional: solo se añade si es_amarilla es True.
    """
    box_l_w = max(175, len(str(l)) * 48 + 30)
    x_resto = 50 + box_l_w + 14
    
    operaciones = [op.strip() for op in str(routing).replace("-", " ").replace(",", " ").split() if op.strip()]
    routing_zpl_lines = []
    y_op_inicio = 430
    espaciado_y = 65
    for idx, op in enumerate(operaciones):
        y_pos = y_op_inicio + (idx * espaciado_y)
        routing_zpl_lines.append(f"^FO525,{y_pos}^A0N,48,60^FD{op}^FS")
    
    bloque_routing = "\n".join(routing_zpl_lines)

    # Bloque inferior condicional para papel amarillo
    bloque_pie = ""
    if es_amarilla:
        bloque_pie = f"""
^FO50,738^GB712,2,2^FS
^FO20,750^FB772,1,0,C,0^A0N,44,55^FD{color_watermark}^FS"""

    zpl = f"""^XA
^CI28
^PW812
^LL812

^FO20,25^FB772,1,0,C,0^A0N,135,120^FD{fecha}^FS

^FO50,165^FB712,2,4,L,0^A0N,36,46^FD{part_name}^FS

^FO50,240^GB{box_l_w},102,2^FS
^FO50,248^FB{box_l_w},1,0,C,0^A0N,68,56^FD{l}^FS
^FO50,316^FB{box_l_w},1,0,C,0^A0N,22,28^FDGRAIN^FS

^FO{x_resto},258^A0N,56,42^FDx  {w}  x  {t}^FS

^FO525,250^A0N,64,78^FDQTY: {qty}^FS

^FO50,375^GB445,75,2^FS
^FO65,395^A0N,36,45^FDTL:^FS

^FO525,375^A0N,36,45^FDROUTING:^FS
{bloque_routing}

^FO50,465^A0N,30,34^FDFace 1:  {face1}^FS
^FO50,510^A0N,30,34^FDFace 2:  {face2}^FS

^FO50,555^A0N,30,34^FDLong 1:^FS^FO175,580^GB320,2,2^FS
^FO50,600^A0N,30,34^FDLong 2:^FS^FO175,625^GB320,2,2^FS
^FO50,645^A0N,30,34^FDShort 1:^FS^FO175,670^GB320,2,2^FS
^FO50,690^A0N,30,34^FDShort 2:^FS^FO175,715^GB320,2,2^FS{bloque_pie}

^XZ"""
    return zpl

def ejecutar_impresion_recuts():
    archivo_excel = obtener_archivo_reciente()
    if not archivo_excel:
        print("Error: No se encontró ningún archivo .xlsx en 'record/' ni en la carpeta raíz.")
        return

    print(f"\n--- Procesando archivo: {archivo_excel} ---")
    wb = openpyxl.load_workbook(archivo_excel, data_only=True)

    if "Catalogo_Lote" not in wb.sheetnames or "Recut_List" not in wb.sheetnames:
        print("Error: Faltan las hojas 'Catalogo_Lote' o 'Recut_List'.")
        return

    # 1. Indexar Catalogo_Lote dando prioridad a PartName_Simplified
    ws_cat = wb["Catalogo_Lote"]
    headers = [str(cell.value).strip() if cell.value is not None else "" for cell in ws_cat[1]]
    col_idx = {h: i for i, h in enumerate(headers)}

    col_unique = col_idx.get("Unique", 0)
    col_name = None
    for h, i in col_idx.items():
        h_limpio = h.upper().replace("\n", "").replace(" ", "")
        if "SIMP" in h_limpio:
            col_name = i
            break
            
    if col_name is None:
        col_name = col_idx.get("PartName", 1)

    col_l = col_idx.get("L", -1)
    col_w = col_idx.get("W", -1)
    col_t = col_idx.get("T", -1)
    col_faces = col_idx.get("Faces", -1)
    col_routing = col_idx.get("Routing", -1)

    catalogo = {}
    for row in ws_cat.iter_rows(min_row=2, values_only=True):
        if not row or row[col_unique] is None:
            continue
        key = str(row[col_unique]).strip()
        catalogo[key] = {
            "PartName": str(row[col_name] or "").strip(),
            "L": str(row[col_l] if col_l != -1 else ""),
            "W": str(row[col_w] if col_w != -1 else ""),
            "T": str(row[col_t] if col_t != -1 else ""),
            "Faces": str(row[col_faces] if col_faces != -1 else ""),
            "Routing": str(row[col_routing] if col_routing != -1 else "RLCUT-KIT")
        }

    # 2. Metadatos (Fecha y Color)
    ws_rec = wb["Recut_List"]
    lote_meta = str(ws_rec["H1"].value or ws_rec["E1"].value or "").strip()
    fecha_etiqueta = formatear_fecha_mes_dia(lote_meta or archivo_excel)
    
    texto_para_color = f"{lote_meta} {archivo_excel}"
    color_detectado = extraer_color_lote(texto_para_color)
    marca_agua = generar_cadena_marca_agua(color_detectado)

    # 3. Barrido de partidas en Recut_List con lectura de Override en Columna C
    recuts_a_imprimir = []
    r = 3
    while True:
        celda_unique = ws_rec.cell(row=r, column=1).value
        celda_qty = ws_rec.cell(row=r, column=2).value
        celda_override_faces = ws_rec.cell(row=r, column=3).value

        if celda_unique is None or str(celda_unique).strip() == "":
            break

        unique_val = str(celda_unique).strip()
        try:
            qty_num = int(float(celda_qty)) if celda_qty is not None else 0
        except ValueError:
            qty_num = 0

        if qty_num > 0:
            if unique_val in catalogo:
                item_data = catalogo[unique_val].copy()
                item_data["Unique"] = unique_val
                item_data["Qty"] = qty_num

                if celda_override_faces is not None and str(celda_override_faces).strip() != "":
                    item_data["Faces"] = str(celda_override_faces).strip()
                    item_data["Override"] = True
                else:
                    item_data["Override"] = False

                recuts_a_imprimir.append(item_data)
            else:
                print(f"[Aviso] Clave '{unique_val}' en fila {r} no existe en catálogo. Omitida.")
        r += 1

    if not recuts_a_imprimir:
        print("\nNo se encontraron recortes registrados con Qty > 0.")
        return

    # 4. Resumen en consola
    print(f"\nFecha en etiqueta: {fecha_etiqueta}")
    print(f"Color de la orden: {color_detectado}")
    print(f"{'#':<3} | {'Unique':<7} | {'Qty':<4} | {'Faces (1/2)':<22} | PartName (Simplificado)")
    print("-" * 85)
    for idx, item in enumerate(recuts_a_imprimir, 1):
        f1, f2 = desglosar_faces_completo(item["Faces"])
        faces_str = f"{f1[:10]}/{f2[:10]}"
        override_mark = "*" if item.get("Override") else " "
        print(f"{idx:<3} | {item['Unique']:<7} | {item['Qty']:<4} | {faces_str:<22}{override_mark}| {item['PartName'][:35]}")

    # 5. Selector de tipo de etiqueta interactivo (y / c)
    es_amarilla = False
    while True:
        tipo_papel = input("\n¿Qué papel tienes cargado en la Zebra? [y = Yellow / c = Color]: ").strip().lower()
        if tipo_papel == 'y':
            es_amarilla = True
            print(f"-> Modo seleccionado: Etiqueta AMARILLA (con marca de agua: '{marca_agua}')")
            break
        elif tipo_papel == 'c':
            es_amarilla = False
            print("-> Modo seleccionado: Etiqueta COLOR DE LA ORDEN (diseño limpio sin pie de página)")
            break
        else:
            print("Opción inválida. Teclea únicamente 'y' para amarilla o 'c' para color.")

    # 6. Confirmación de envío
    confirmacion = input(f"\n¿Enviar {len(recuts_a_imprimir)} etiqueta(s) a la Zebra ZT230 ({IP_ZEBRA})? (s/n): ").strip().lower()
    if confirmacion != 's':
        print("Impresión cancelada.")
        return

    # 7. Envío RAW a la impresora
    try:
        with socket.socket(socket.AF_INET, socket.SOCK_STREAM) as s:
            s.settimeout(6.0)
            s.connect((IP_ZEBRA, PUERTO))
            for item in recuts_a_imprimir:
                f1, f2 = desglosar_faces_completo(item["Faces"])
                zpl = generar_zpl_maestro(
                    fecha=fecha_etiqueta,
                    part_name=item["PartName"],
                    l=item["L"],
                    w=item["W"],
                    t=item["T"],
                    qty=item["Qty"],
                    routing=item["Routing"],
                    face1=f1,
                    face2=f2,
                    es_amarilla=es_amarilla,
                    color_watermark=marca_agua
                )
                s.sendall(zpl.encode("utf-8"))
        print(f"\n¡Éxito! Se imprimieron {len(recuts_a_imprimir)} etiqueta(s) correctamente.")
    except Exception as e:
        print(f"Error de comunicación con la impresora: {e}")

if __name__ == "__main__":
    ejecutar_impresion_recuts()